import json,re,html,sys
from pathlib import Path
from html.parser import HTMLParser
from collections import defaultdict
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
class P(HTMLParser):
 def __init__(self):super().__init__(convert_charrefs=True);self.out=[]
 def handle_data(self,x):self.out.append(x)
 def handle_starttag(self,tag,a):
  if tag in ('p','br'):self.out.append('\n')
def clean(s):
 if '<' not in s:return html.unescape(s)
 p=P();p.feed(s);return ''.join(p.out).strip()
def category(r):
 p=r['path'];f=r['file']
 if p.startswith('labels.'):return '01 · 通用界面与语言切换'
 if p.startswith(('ambientMeta','calendarIntro','cycleNames')) or f=='calendar-motion.js':return '02 · 首页日名与 Calendar'
 if p.startswith('works['):return '03 · Works / 作品'
 if p.startswith('collections[0]'):return '04 · Collection / 墟域图与遗构馆'
 if p.startswith('collections[1]'):return '05 · Collection / room by the lake'
 if p.startswith(('seawater','collections[2]')):return '06 · Collection / collection of (all) oceans and seas'
 if p.startswith(('collections[3]','fragrance','theme')):return '07 · Collection / 芬芳厅'
 if p.startswith('writings['):return '08 · Writings / 文稿'
 if f.startswith('ruin-mini'):return '09 · 地图预览与辅助文字'
 return '10 · 其他内容'
data=json.loads(Path('.editorial/site-copy-map.json').read_text(encoding='utf8'))
d=Document(); sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(2);sec.bottom_margin=Cm(2);sec.left_margin=Cm(2.3);sec.right_margin=Cm(2.3)
normal=d.styles['Normal'];normal.font.name='Noto Serif CJK SC';normal.font.size=Pt(10.5);normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC');normal.paragraph_format.space_after=Pt(4)
for name,size in [('Title',25),('Heading 1',17),('Heading 2',12)]:
 st=d.styles[name];st.font.name='Noto Serif CJK SC';st.font.size=Pt(size);st._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC')
d.add_heading('全站文字校订稿',0);d.add_paragraph('sky-sea-lake-cloud  ·  中文 / English / 日本語')
d.add_paragraph('编辑说明',style='Heading 1')
for text in ['这是从当前 GitHub 版本提取的网页文字：作品名称、介绍、图注、长篇笔记、菜单、日历和图表文案。','请只修改中文／EN／日本語文字栏；保留 TXT 编号，后续可以按编号替换回网页。','不需更改的版本原样保留。要删除某段，请填写【删除此段】，不要删除编号。','带链接的正文在此仅显示可读文字；重新导入时会核对链接。外部独立网站、图片中写死的文字和程序变量不在此稿内。']:
 d.add_paragraph('— '+text)
d.add_paragraph('源码版本：'+str(data.get('revision') or 'main'))
groups=defaultdict(list)
for item in data['records']:groups[category(item)].append(item)
d.add_page_break();d.add_heading('目录',level=1)
for group,rs in sorted(groups.items()):d.add_paragraph(group+'  ·  '+str(len(rs))+' 组')
for group,rs in sorted(groups.items()):
 d.add_page_break();d.add_heading(group,level=1)
 d.add_paragraph(str(len(rs))+' 组可编辑文字')
 for item in rs:
  title=item['id']+'  ·  '+item['path']
  d.add_heading(title,level=2)
  p=d.add_paragraph(f"{item['file']}  /  L{item['line']}")
  p.runs[0].font.size=Pt(8);p.runs[0].font.color.rgb=RGBColor(135,140,138)
  table=d.add_table(rows=3,cols=2);table.autofit=False
  for i,(lang,text) in enumerate([('中文',item['languages'].get('zh','')),('EN',item['languages'].get('en','')),('日本語',item['languages'].get('ja',''))]):
   a,b=table.rows[i].cells;a.width=Cm(1.7);b.width=Cm(14.5)
   a.text=lang;b.text=clean(text) or ' '
   for cell,c in [(a,'F0F0EB'),(b,'FAFAF7')]:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),c);cell._tc.get_or_add_tcPr().append(shade)
  d.add_paragraph(' ')
d.core_properties.title='sky-sea-lake-cloud 全站文字校订稿'
out=Path('downloads/site-copy-editable.docx');out.parent.mkdir(exist_ok=True);d.save(out)
test=Document(out);print('DOCX saved:',out,'groups:',len(data['records']),'tables:',len(test.tables))
assert len(test.tables)==len(data['records'])
